from __future__ import annotations

import os


def create_excel_sheet(file_name: str, headers: list, rows: list | None = None) -> str:
    from openpyxl import Workbook

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Sheet1"
    worksheet.append(headers)
    for row in rows or []:
        worksheet.append(row)
    file_path = os.path.abspath(file_name if file_name.endswith(".xlsx") else f"{file_name}.xlsx")
    workbook.save(file_path)
    return f"Success: Excel sheet created at {file_path}"


def create_word_document(file_name: str, title: str, content: str) -> str:
    from docx import Document

    document = Document()
    document.add_heading(title, level=1)
    document.add_paragraph(content)
    file_path = os.path.abspath(file_name if file_name.endswith(".docx") else f"{file_name}.docx")
    document.save(file_path)
    return f"Success: Word document created at {file_path}"